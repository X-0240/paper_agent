import argparse
import json
import os
import random
import re
import sys

sys.path.insert(0,os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font

from evaluation.metrics import norm

BASE=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SECTIONS_DIR=os.path.join(BASE,"datasets","v300","sections")

#openpyxl 拒绝控制字符，PDF 解析文本里会混入，写表前统一清洗
ILLEGAL=re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


def sanitize(value,limit=30000):
    text=value if isinstance(value,str) else str(value)
    text=ILLEGAL.sub("",text)
    return text[:limit]

#按四类命中结果分层抽样：小桶全抽，大桶少量对照
QUOTA={
    "仅基线命中":10,
    "仅候选命中":18,
    "两边都没中":14,
    "两边都中":8,
}


def load_json(path):
    with open(path,encoding="utf-8") as f:
        return json.load(f)


def locate(sections,gold):
    #在章节缓存里定位金标句，返回(章节标题, 章节文本, 匹配位置)
    target=norm(gold)
    if not target:
        return "","",-1
    for sec in sections:
        body=sec.get("text","")
        pos=norm(body).find(target)
        if pos>=0:
            return sec.get("title",""),body,pos
    return "","",-1


def context_window(body,pos,span=700):
    #原文按字符截取上下文，供人工核对金标句是否真实存在
    start=max(0,pos-span)
    end=min(len(body),pos+span)
    head="…" if start>0 else ""
    tail="…" if end<len(body) else ""
    return f"{head}{body[start:end]}{tail}"


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--sealed",default=os.path.join(BASE,"evaluation","results","v300_sealed_test.json"))
    parser.add_argument("--questions",default=os.path.join(BASE,"evaluation","questions","v300_test_reviewed.json"))
    parser.add_argument("--output",default=os.path.join(BASE,"evaluation","results","v300_label_audit.xlsx"))
    parser.add_argument("--sample-output",default=os.path.join(BASE,"evaluation","results","v300_label_audit_sample.json"))
    parser.add_argument("--seed",type=int,default=20260918)
    args=parser.parse_args()

    sealed=load_json(args.sealed)
    questions={q["question"].strip():q for q in load_json(args.questions)}
    rng=random.Random(args.seed)

    strata={key:[] for key in QUOTA}
    for row in sealed["rows"]:
        question=row["question"].strip()
        if row["baseline_hit5"] and row["candidate_hit5"]:
            bucket="两边都中"
        elif row["baseline_hit5"]:
            bucket="仅基线命中"
        elif row["candidate_hit5"]:
            bucket="仅候选命中"
        else:
            bucket="两边都没中"
        strata[bucket].append((question,row))

    picked=[]
    for bucket,quota in QUOTA.items():
        items=strata[bucket]
        take=items if len(items)<=quota else rng.sample(items,quota)
        take=sorted(take,key=lambda x:x[0])
        for question,row in take:
            picked.append((bucket,question,row))
    print("分层抽样：")
    for bucket in QUOTA:
        print(f"  {bucket}: 总体 {len(strata[bucket])}，抽 {sum(1 for p in picked if p[0]==bucket)}")

    cache={}
    wb=Workbook()
    ws=wb.active
    ws.title="证据抽检"
    header=["序号","桶","论文","问题","金标证据句","标注章节","定位章节","原文上下文","能对上原文?(是/否)","错误类型","备注"]
    ws.append(header)
    for cell in ws[1]:
        cell.font=Font(bold=True)
    auto_found=0
    for idx,(bucket,question,row) in enumerate(picked,1):
        item=questions.get(question)
        if item is None:
            continue
        paper=item.get("paper_id","")
        if paper not in cache:
            path=os.path.join(SECTIONS_DIR,f"{paper}.json")
            cache[paper]=load_json(path) if os.path.exists(path) else []
        sections=cache[paper]
        golds=item.get("evidence_sentences") or []
        gold_text="\n\n".join(f"[{i}] {g}" for i,g in enumerate(golds,1))
        titles=[];ctxs=[];found=False
        for g in golds:
            title,body,pos=locate(sections,g)
            if pos>=0:
                found=True
                titles.append(title)
                ctxs.append(f"【{title}】\n{context_window(body,pos)}")
            else:
                ctxs.append("【未在章节缓存里定位到该句，需人工确认】")
        if found:
            auto_found+=1
        ws.append([sanitize(x) if isinstance(x,str) else x for x in [
            idx,bucket,paper,question,gold_text,
            item.get("section","")," / ".join(t for t in titles if t),
            "\n\n".join(ctxs),"","","",
        ]])
    widths=[6,12,22,46,60,24,24,80,16,16,20]
    for col,w in enumerate(widths,1):
        ws.column_dimensions[ws.cell(row=1,column=col).column_letter].width=w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,str):
                cell.value=sanitize(cell.value)
            cell.alignment=Alignment(wrap_text=True,vertical="top")

    guide=wb.create_sheet("填写说明")
    for line in [
        "判定方法：只看'金标证据句'能否在该论文原文里找到，并且能否支撑问题的答案。",
        "只看原文，不要借助外部常识；原文里找不到、或找到但不支持答案，都算'否'。",
        "'错误类型'可填：金标句不在原文 / 原文支持弱 / 章节错 / 论文错 / 其他。",
        "'备注'写你的依据，例如原文实际写的是什么。",
        "分层抽样结论只能代表抽检的这 50 题，不能推广成整体准确率。",
        f"本表共 {len(picked)} 题；脚本自动定位成功的题数会在控制台打印，用于和人工判定对照。",
    ]:
        guide.append([sanitize(line)])
    guide.column_dimensions["A"].width=100
    os.makedirs(os.path.dirname(args.output),exist_ok=True)
    wb.save(args.output)
    with open(args.sample_output,"w",encoding="utf-8") as f:
        json.dump({
            "seed":args.seed,
            "quota":QUOTA,
            "strata_total":{k:len(v) for k,v in strata.items()},
            "sample":[{"bucket":b,"question":q} for b,q,_ in picked],
        },f,ensure_ascii=False,indent=1)
    print(f"自动定位成功：{auto_found}/{len(picked)}")
    print(f"已保存：{args.output}")
    print(f"抽样清单：{args.sample_output}")


if __name__=="__main__":
    main()
